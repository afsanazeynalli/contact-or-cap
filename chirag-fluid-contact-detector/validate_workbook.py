from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
BOOK = ROOT / "Fluid_Contact_Detector_CHIRAG.xlsx"

required_sheets = {
    "Dashboard",
    "Master_Data",
    "Settings",
    "Contact_Detector",
    "Contact_Summary",
    "Methodology",
}

if not BOOK.exists():
    raise FileNotFoundError(BOOK)

wb = load_workbook(BOOK, read_only=True, data_only=False)

missing = required_sheets.difference(wb.sheetnames)
if missing:
    raise RuntimeError(f"Missing sheets: {sorted(missing)}")

headers = [c.value for c in next(wb["Master_Data"].iter_rows(min_row=1, max_row=1))]
required_headers = {
    "Well", "Formation", "TVD-SCS", "Fluids",
    "TST_S (optional)", "Fluid_Override (optional)", "Fluid_Class"
}
missing_headers = required_headers.difference(headers)

if missing_headers:
    raise RuntimeError(f"Missing Master_Data headers: {sorted(missing_headers)}")

print("Workbook structure validation passed.")
